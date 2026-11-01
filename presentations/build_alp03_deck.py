#!/usr/bin/env python3
"""Пробная презентация модуля 3 «Лавины и влияние высокогорной среды».

Слайды разбирают учебную оболочку комплекта: четыре урока, задания
и контрольный тест. Страницы пособия остаются источником рисунков
и определений. Верные формулировки теста взяты из ключа модуля.
"""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

W = Inches(13.333)
H = Inches(7.5)

NAVY = RGBColor(0x0B, 0x2F, 0x4A)
BLUE = RGBColor(0x14, 0x66, 0xA8)
ICE = RGBColor(0xE7, 0xF0, 0xF7)
BG = RGBColor(0xF6, 0xF8, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x12, 0x16, 0x1B)
MUTED = RGBColor(0x4A, 0x55, 0x60)
RED = RGBColor(0xA3, 0x2D, 0x2D)
RED_SOFT = RGBColor(0xF8, 0xEE, 0xEC)
LINE = RGBColor(0xD5, 0xDE, 0xE6)
GOLD = RGBColor(0x8A, 0x6A, 0x12)
GOLD_SOFT = RGBColor(0xF8, 0xF3, 0xE3)

FONT = "Arial"
OUT = "/workspace/presentations/ALP-03_laviny-i-vysota_proba.pptx"


def _font(run, name, size, color, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rpr, qn(f"a:{tag}"))
        el.set("typeface", name)


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    _fill(shape, color)
    shape.shadow.inherit = False
    return shape


def _box(slide, x, y, w, h, color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = 0.08
    if color is not None:
        _fill(shape, color)
    shape.shadow.inherit = False
    return shape


def _text(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(x, y, w, h)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        _font(run, FONT, size, color, bold)
    return shape


def _notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def _footer(slide, n, total, light=False):
    color = RGBColor(0xC5, 0xD4, 0xE2) if light else MUTED
    _text(
        slide,
        Inches(0.55),
        Inches(7.12),
        Inches(10.2),
        Inches(0.28),
        "Школа альпинизма  ·  Модуль 3  ·  пробная презентация  ·  не аттестация",
        12,
        color,
    )
    _text(
        slide,
        Inches(11.4),
        Inches(7.12),
        Inches(1.4),
        Inches(0.28),
        f"{n}  /  {total}",
        12,
        color,
        align=PP_ALIGN.RIGHT,
    )


def _bg(slide, color):
    _rect(slide, 0, 0, W, H, color)


def _kicker(slide, text, y=Inches(0.38), color=BLUE):
    _text(slide, Inches(0.55), y, Inches(12.2), Inches(0.28), text.upper(), 13, color, bold=True)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = W
        self.prs.slide_height = H
        self.prs.core_properties.title = (
            "Школа альпинизма. Модуль 3. Лавины и влияние высокогорной среды"
        )
        self.prs.core_properties.subject = "Пробная презентация для теста формата"
        self.prs.core_properties.author = "Учебный комплект курса по пособию ФАР"
        self.slides = []

    def add(self, kind, **kwargs):
        self.slides.append((kind, kwargs))

    def build(self):
        total = len(self.slides)
        for i, (kind, kw) in enumerate(self.slides, start=1):
            layout = self.prs.slide_layouts[6]
            slide = self.prs.slides.add_slide(layout)
            getattr(self, kind)(slide, i, total, **kw)
        self.prs.save(OUT)

    def title(self, slide, n, total):
        _bg(slide, NAVY)
        _rect(slide, 0, 0, Inches(0.18), H, BLUE)
        _text(slide, Inches(0.7), Inches(1.35), Inches(11.5), Inches(0.32),
              "ШКОЛА АЛЬПИНИЗМА  ·  ОНЛАЙН-КУРС", 15, RGBColor(0x9E, 0xC9, 0xEC), bold=True)
        _text(slide, Inches(0.7), Inches(1.85), Inches(11.8), Inches(0.4),
              "МОДУЛЬ 3", 18, WHITE, bold=True)
        _text(slide, Inches(0.7), Inches(2.35), Inches(11.8), Inches(1.8),
              "Лавины и влияние\nвысокогорной среды", 48, WHITE, bold=True)
        _text(slide, Inches(0.7), Inches(4.55), Inches(10.5), Inches(0.7),
              "Пробная презентация: разбор пособия на слайды и контрольный тест.",
              20, RGBColor(0xD5, 0xE4, 0xF0))
        _text(slide, Inches(0.7), Inches(5.7), Inches(11), Inches(0.7),
              "4–5 часов теории  ·  пособие, с. 104–139\nНе допуск к действиям, не квалификация и не аттестация инструктора-проводника",
              16, RGBColor(0x9E, 0xC9, 0xEC))
        _notes(slide, "Первый готовый модуль комплекта. История и природа гор отдельным пособием в эту поставку не входили, поэтому пробный формат собран здесь.")

    def bullets(self, slide, n, total, kicker, title, items, note="", tone="plain"):
        _bg(slide, BG if tone == "plain" else WHITE)
        _rect(slide, 0, 0, W, Inches(0.08), BLUE if tone != "danger" else RED)
        _kicker(slide, kicker, color=RED if tone == "danger" else BLUE)
        _text(slide, Inches(0.55), Inches(0.72), Inches(12.2), Inches(0.9), title, 32, INK, bold=True)
        top = Inches(1.85)
        for item in items:
            _rect(slide, Inches(0.58), top + Inches(0.1), Inches(0.08), Inches(0.08), BLUE if tone != "danger" else RED)
            _text(slide, Inches(0.9), top, Inches(11.7), Inches(0.72), item, 22, INK)
            top += Inches(0.82)
        _footer(slide, n, total)
        _notes(slide, note)

    def cards(self, slide, n, total, kicker, title, cards, note="", columns=2):
        _bg(slide, BG)
        _rect(slide, 0, 0, W, Inches(0.08), BLUE)
        _kicker(slide, kicker)
        _text(slide, Inches(0.55), Inches(0.7), Inches(12.2), Inches(0.85), title, 30, INK, bold=True)
        gap = Inches(0.22)
        left = Inches(0.55)
        width = Inches(12.25)
        count = len(cards)
        cols = columns
        rows = (count + cols - 1) // cols
        card_w = (width - gap * (cols - 1)) / cols
        card_h = Inches(1.7) if rows > 1 else Inches(2.15)
        top0 = Inches(1.9)
        for i, card in enumerate(cards):
            r, c = divmod(i, cols)
            x = left + (card_w + gap) * c
            y = top0 + (card_h + gap) * r
            _box(slide, x, y, card_w, card_h, WHITE)
            _rect(slide, x, y, Inches(0.08), card_h, BLUE)
            _text(slide, x + Inches(0.28), y + Inches(0.22), card_w - Inches(0.45), Inches(0.4),
                  card["label"], 14, BLUE, bold=True)
            _text(slide, x + Inches(0.28), y + Inches(0.68), card_w - Inches(0.45), card_h - Inches(0.85),
                  card["text"], 18, INK)
        _footer(slide, n, total)
        _notes(slide, note)

    def zones(self, slide, n, total):
        _bg(slide, BG)
        _rect(slide, 0, 0, W, Inches(0.08), BLUE)
        _kicker(slide, "Урок 1  ·  пособие, с. 104–109")
        _text(slide, Inches(0.55), Inches(0.7), Inches(12.2), Inches(0.8),
              "Три зоны, которые нужно показать на склоне", 30, INK, bold=True)
        bands = [
            (NAVY, "1", "Зона зарождения", "Где снег может прийти в движение"),
            (BLUE, "2", "Зона транзита", "Где лавина набирает путь"),
            (RGBColor(0x3E, 0x6F, 0x96), "3", "Зона отложения", "Где масса останавливается"),
        ]
        y = Inches(1.9)
        for color, num, name, desc in bands:
            _box(slide, Inches(0.55), y, Inches(7.5), Inches(1.35), color)
            _text(slide, Inches(0.8), y + Inches(0.28), Inches(1.0), Inches(0.7), num, 32, WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
            _text(slide, Inches(1.9), y + Inches(0.22), Inches(5.7), Inches(0.45), name, 22, WHITE, bold=True)
            _text(slide, Inches(1.9), y + Inches(0.72), Inches(5.7), Inches(0.4), desc, 16, RGBColor(0xE7, 0xF0, 0xF7))
            y += Inches(1.5)
        _box(slide, Inches(8.3), Inches(1.9), Inches(4.5), Inches(4.35), WHITE)
        _text(slide, Inches(8.55), Inches(2.15), Inches(4.05), Inches(0.4), "ВОПРОС ГРУППЕ", 13, BLUE, bold=True)
        _text(slide, Inches(8.55), Inches(2.65), Inches(4.05), Inches(3.2),
              "В какой зоне мы окажемся и что из этого следует для линии движения?\n\nСхема учебная. Рисунок склона в пособии не заменять этим слайдом.",
              18, INK)
        _footer(slide, n, total)
        _notes(slide, "Попросить нарисовать три зоны до обсуждения «идти или нет». Пока зоны не названы, решение рано.")

    def split(self, slide, n, total, kicker, title, left, right, note=""):
        _bg(slide, BG)
        _rect(slide, 0, 0, W, Inches(0.08), BLUE)
        _kicker(slide, kicker)
        _text(slide, Inches(0.55), Inches(0.7), Inches(12.2), Inches(0.8), title, 30, INK, bold=True)
        _box(slide, Inches(0.55), Inches(1.9), Inches(6.0), Inches(4.7), WHITE)
        _rect(slide, Inches(0.55), Inches(1.9), Inches(6.0), Inches(0.1), BLUE)
        _text(slide, Inches(0.85), Inches(2.2), Inches(5.4), Inches(0.4), left["label"], 14, BLUE, bold=True)
        _text(slide, Inches(0.85), Inches(2.75), Inches(5.4), Inches(3.4), left["text"], 20, INK)
        _box(slide, Inches(6.8), Inches(1.9), Inches(6.0), Inches(4.7), WHITE)
        _rect(slide, Inches(6.8), Inches(1.9), Inches(6.0), Inches(0.1), RED)
        _text(slide, Inches(7.1), Inches(2.2), Inches(5.4), Inches(0.4), right["label"], 14, RED, bold=True)
        _text(slide, Inches(7.1), Inches(2.75), Inches(5.4), Inches(3.4), right["text"], 20, INK)
        _footer(slide, n, total)
        _notes(slide, note)

    def section(self, slide, n, total, number, title, pages, note=""):
        _bg(slide, NAVY)
        _text(slide, Inches(0.7), Inches(2.05), Inches(11), Inches(0.4), number, 16, RGBColor(0x9E, 0xC9, 0xEC), bold=True)
        _text(slide, Inches(0.7), Inches(2.55), Inches(11.5), Inches(1.6), title, 42, WHITE, bold=True)
        _text(slide, Inches(0.7), Inches(4.5), Inches(11), Inches(0.4), pages, 18, RGBColor(0xD5, 0xE4, 0xF0))
        _footer(slide, n, total, light=True)
        _notes(slide, note)

    def task(self, slide, n, total, kicker, title, beginner, instructor, note=""):
        _bg(slide, BG)
        _rect(slide, 0, 0, W, Inches(0.08), BLUE)
        _kicker(slide, kicker)
        _text(slide, Inches(0.55), Inches(0.7), Inches(12.2), Inches(0.7), title, 30, INK, bold=True)
        _box(slide, Inches(0.55), Inches(1.75), Inches(6.0), Inches(4.85), WHITE)
        _text(slide, Inches(0.85), Inches(2.0), Inches(5.4), Inches(0.35), "НАЧИНАЮЩИЙ", 13, BLUE, bold=True)
        _text(slide, Inches(0.85), Inches(2.5), Inches(5.4), Inches(3.7), beginner, 20, INK)
        _box(slide, Inches(6.8), Inches(1.75), Inches(6.0), Inches(4.85), WHITE)
        _text(slide, Inches(7.1), Inches(2.0), Inches(5.4), Inches(0.35), "МЕТОДИЧЕСКАЯ ПРАКТИКА", 13, BLUE, bold=True)
        _text(slide, Inches(7.1), Inches(2.5), Inches(5.4), Inches(3.7), instructor, 20, INK)
        _footer(slide, n, total)
        _notes(slide, note)

    def question(self, slide, n, total, qn, prompt, options, note=""):
        _bg(slide, BG)
        _rect(slide, 0, 0, W, Inches(0.08), GOLD)
        _kicker(slide, f"Контрольный тест  ·  вопрос {qn} из 6", color=GOLD)
        _text(slide, Inches(0.55), Inches(0.75), Inches(12.2), Inches(1.35), prompt, 30, INK, bold=True)
        y = Inches(2.35)
        for letter, text in options:
            _box(slide, Inches(0.55), y, Inches(12.25), Inches(1.25), WHITE)
            bubble = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), y + Inches(0.32), Inches(0.58), Inches(0.58))
            _fill(bubble, ICE)
            bubble.shadow.inherit = False
            _text(slide, Inches(0.8), y + Inches(0.4), Inches(0.58), Inches(0.42), letter, 18, BLUE, bold=True, align=PP_ALIGN.CENTER)
            _text(slide, Inches(1.65), y + Inches(0.32), Inches(10.8), Inches(0.7), text, 22, INK, anchor=MSO_ANCHOR.MIDDLE)
            y += Inches(1.4)
        _footer(slide, n, total)
        _notes(slide, note)

    def key(self, slide, n, total):
        _bg(slide, NAVY)
        _text(slide, Inches(0.55), Inches(0.38), Inches(12), Inches(0.3),
              "КЛЮЧ ПРЕПОДАВАТЕЛЯ  ·  НЕ ПОКАЗЫВАТЬ ДО ОТВЕТОВ", 14, RGBColor(0xE7, 0xC8, 0xC4), bold=True)
        _text(slide, Inches(0.55), Inches(0.75), Inches(12), Inches(0.55),
              "Верные формулировки из ключа модуля", 28, WHITE, bold=True)
        rows = [
            ("1  Б", "Снег, погода, рельеф и воздействие на людей. Оценка комплексная."),
            ("2  Б", "След предыдущей группы — недостаточная информация, не гарантия."),
            ("3  А", "Красный флаг — повод пересмотреть план, а не разрешение идти."),
            ("4  Б", "Онлайн-тест не подтверждает навык поиска. Нужна очная проверка."),
            ("5  Б", "Оценить состояние и пересмотреть план, при необходимости запросить помощь."),
            ("6  А", "Личная реакция на высоту может отличаться от формы и от плана."),
        ]
        y = Inches(1.55)
        for label, text in rows:
            _text(slide, Inches(0.55), y, Inches(1.3), Inches(0.7), label, 20, RGBColor(0xFF, 0xCF, 0x3F), bold=True)
            _text(slide, Inches(1.9), y, Inches(10.6), Inches(0.7), text, 18, WHITE)
            y += Inches(0.78)
        _footer(slide, n, total, light=True)
        _notes(slide, "Буквы и верные формулировки совпадают с ключом модуля. Варианты-ловушки на слайдах вопросов собраны из ложных выводов самого модуля: в рабочей копии PDF полный перечень отвлекающих ответов на странице теста сохранился не целиком. Вопросы 2, 3, 4 и 5 — критические: ошибка в них не закрывается общим баллом. Это проектное правило курса, не норматив ФАР.")

    def result(self, slide, n, total):
        _bg(slide, BG)
        _rect(slide, 0, 0, W, Inches(0.08), BLUE)
        _kicker(slide, "После теста")
        _text(slide, Inches(0.55), Inches(0.7), Inches(12.2), Inches(0.7),
              "Как читать шесть ответов", 32, INK, bold=True)
        cards = [
            ("Порог", "Пять верных из шести — ориентир 80%. Это порог курса, не балл экзамена ФАР."),
            ("Критические", "Ошибка в вопросах 2, 3, 4 или 5 оставляет модуль незакрытым при любом общем счёте."),
            ("Два статуса", "«Тест пройден» и «навык отработан очно» ставятся отдельно. Второй этот слайд не даёт."),
        ]
        x = Inches(0.55)
        for label, text in cards:
            _box(slide, x, Inches(1.8), Inches(3.95), Inches(3.3), WHITE)
            _text(slide, x + Inches(0.28), Inches(2.05), Inches(3.4), Inches(0.4), label, 16, BLUE, bold=True)
            _text(slide, x + Inches(0.28), Inches(2.6), Inches(3.4), Inches(2.1), text, 18, INK)
            x += Inches(4.15)
        _text(slide, Inches(0.55), Inches(5.35), Inches(12.2), Inches(1.4),
              "Дальше — страницы пособия 104–139, без замены текста. Ответы записывают до просмотра ключа.\nПоиск, зондирование, откапывание, первая помощь и движение по снегу остаются очной отработкой.",
              18, MUTED)
        _footer(slide, n, total)
        _notes(slide, "Если пилот покажет, что слушатели путают тест с допуском, на следующем модуле вынести эту рамку ещё и устным вопросом в начале.")


def main():
    d = Deck()
    d.add("title")
    d.add(
        "bullets",
        kicker="Как пользоваться",
        title="Пособие остаётся источником. Слайды ведут занятие.",
        items=[
            "Сначала слайд и разговор, затем сверка с рисунком и определением в пособии.",
            "Задания выполняются до ключа. Ключ в конце — только для преподавателя.",
            "Это проба формата: если занятие читается, так же разбираем остальные модули.",
            "Презентация не заменяет 4–5 часов работы с страницами 104–139.",
        ],
        note="Не читать слайды подряд. После каждой пары тезисов возвращать к странице пособия.",
    )
    d.add(
        "cards",
        kicker="Два режима одного модуля",
        title="Один комплект, разная работа",
        cards=[
            {"label": "НАЧИНАЮЩИЙ", "text": "Пройти четыре урока, сделать схемы и сдать шесть вопросов. Полевой навык возникает только с инструктором."},
            {"label": "КАНДИДАТ В ИНСТРУКТОРЫ", "text": "Дополнительно объяснить ограничение метода, разобрать решение группы и назвать критерий очного зачёта."},
        ],
        note="Материал не воспроизводит и не утверждает действующие условия аттестации ФАР.",
    )
    d.add(
        "bullets",
        kicker="Граница безопасности",
        title="Слайд объясняет решение. Он не выдаёт допуск.",
        items=[
            "Можно разобрать: где опасность, чего не хватает в данных, когда отказаться.",
            "Нельзя зачесть: поиск в лавине, зондирование, откапывание, первую помощь.",
            "Нельзя зачесть: страховку, спуск, снежную и ледовую технику.",
            "Статус комплекта: рабочий материал, не свидетельство о квалификации.",
        ],
        tone="danger",
        note="Произнести границу вслух до урока 1. Иначе тест начинают воспринимать как допуск.",
    )
    d.add(
        "cards",
        kicker="Маршрут модуля",
        title="Четыре урока и один тест",
        columns=2,
        cards=[
            {"label": "УРОК 1", "text": "Механизм и классификации. Пособие, с. 104–109."},
            {"label": "УРОК 2", "text": "Погода, рельеф и тревожные признаки. С. 107–118."},
            {"label": "УРОК 3", "text": "Поиск: только рамка теории. С. 119–121."},
            {"label": "УРОК 4", "text": "Холод, солнце, высота, питание и вода. С. 122–139."},
        ],
        note="После четырёх уроков — ситуация «знакомый склон после снегопада и ветра» и шесть вопросов.",
    )
    d.add("section", number="УРОК 1", title="Механизм\nи классификации", pages="Пособие, с. 104–109",
          note="Цель: объяснить механизм, видеть предел оценки конкретного склона, отделять наблюдение от вывода.")
    d.add(
        "bullets",
        kicker="Урок 1",
        title="Риск — это не название типа лавины",
        items=[
            "Он складывается из снега, погоды, рельефа и того, окажется ли там человек.",
            "Большая часть инцидентов с людьми связана с их собственными решениями.",
            "Снег лежит слоями. Слои копились в разные периоды и ведут себя по-разному.",
            "«Доска» опасна тем, что сразу приходит в движение большой объём снега.",
        ],
        note="Не устраивать классификацию ради классификации. Вопрос урока: где путь лавины и где группа.",
    )
    d.add("zones")
    d.add(
        "cards",
        kicker="Урок 1",
        title="Это наблюдение. Это не гарантия склона.",
        cards=[
            {"label": "СВОДКА ПО РАЙОНУ", "text": "Общая оценка не отвечает за конкретный склон и конкретный час."},
            {"label": "НЕТ СТАРЫХ СЛЕДОВ", "text": "Отсутствие следов лавин не доказывает, что схода не будет."},
            {"label": "ЧУЖОЙ СЛЕД", "text": "Кто-то прошёл раньше — это факт прохода, не проверка устойчивости."},
            {"label": "«Я ЗДЕСЬ ХОДИЛ»", "text": "Память о маршруте не заменяет сегодняшний снег, ветер и состав группы."},
        ],
        note="Спросить: какое из четырёх группа чаще всего использует как разрешение идти.",
    )
    d.add(
        "task",
        kicker="Урок 1  ·  задание",
        title="Сначала схема, потом вывод",
        beginner="Нарисуйте склон: где зарождение, путь и отложение. Подпишите три вопроса, на которые нужно ответить до выхода.",
        instructor="Попросите отделить факты от интерпретаций. Где не хватает данных и почему опыт не может эти данные «добавить».",
        note="Самопроверка вслух: что известно, что неизвестно, какой признак заставит остановиться.",
    )
    d.add("section", number="УРОК 2", title="Погода, рельеф,\nтревожные признаки", pages="Пособие, с. 107–118",
          note="Цель: перечислить факторы, работать с неопределённостью, предложить смену маршрута.")
    d.add(
        "cards",
        kicker="Урок 2",
        title="Что меняет решение по склону",
        cards=[
            {"label": "ОСАДКИ И ВЕТЕР", "text": "Свежий снег и перенос на подветренную сторону копят массу там, где её не ждали. Карниз — часть этой картины."},
            {"label": "ТЕМПЕРАТУРА", "text": "Важен ход температуры, а не одна цифра утром. Пособие разбирает, как он меняет устойчивость."},
            {"label": "ФОРМА И ЭКСПОЗИЦИЯ", "text": "Крутизна, форма, сторона света и местные накопления решают больше, чем «район в целом»."},
            {"label": "ПОСЛЕДСТВИЯ", "text": "Даже редкий сход имеет значение, если под ним стоит группа. Схему углов смотреть в пособии, с. 111."},
        ],
        note="Не воспроизводить по памяти таблицу углов со страницы 111: на слайде её нет специально.",
    )
    d.add(
        "split",
        kicker="Урок 2",
        title="Красный флаг — не выключатель риска",
        left={"label": "ЧЕМ ОН ЯВЛЯЕТСЯ", "text": "Поводом пересмотреть план.\n\nЧастью общей картины вместе с погодой, рельефом и последствиями."},
        right={"label": "ЧЕМ ОН НЕ ЯВЛЯЕТСЯ", "text": "Ответом «можно».\n\nОтсутствие флага не делает склон безопасным. Тест на снегу не чинит плохую линию."},
        note="Формулировка теста: красные флаги — повод пересмотреть план.",
    )
    d.add(
        "bullets",
        kicker="Урок 2",
        title="Знакомый маршрут после ветра — уже другой",
        items=[
            "Снегопад и ветер меняют склон, даже если линия на карте та же.",
            "Нужна менее опасная альтернатива и заранее названная точка отказа.",
            "Скорость прохождения не превращает плохую линию в рабочую.",
            "Решение «не идти» — такой же результат урока, как верный термин.",
        ],
        note="Связать с итоговой ситуацией теста, но ответы ещё не обсуждать.",
    )
    d.add(
        "task",
        kicker="Урок 2  ·  задание",
        title="Разобрать признаки, не маршрут целиком",
        beginner="Из описания выпишите погодные и рельефные признаки. Разделите: известно, неизвестно, требует немедленного пересмотра.",
        instructor="Дайте два варианта линии. Проведите разговор, почему от одного группа отказывается.",
        note="Следить, чтобы спор не свёлся к тому, кто смелее.",
    )
    d.add("section", number="УРОК 3", title="Поиск в лавине:\nрамка, не навык", pages="Пособие, с. 119–121",
          note="Время до прибытия спасателей критично. Это причина знать рамку. Это не причина считать просмотр слайдов тренировкой.")
    d.add(
        "split",
        kicker="Урок 3",
        title="Что слайд может и чего не может",
        left={"label": "МОЖНО ПРОВЕРИТЬ ЗДЕСЬ", "text": "Понимает ли человек, что группа действует до прибытия помощи и что без навыка на этот рельеф не выходят."},
        right={"label": "НЕЛЬЗЯ ЗАЧЕСТЬ ЗДЕСЬ", "text": "Поиск, зондирование, откапывание и организацию помощи. Это регулярная очная тренировка с проверкой действий."},
        note="Не показывать пошаговую технику поиска. В пособии она есть для очного курса, не для самостоятельного применения.",
    )
    d.add(
        "bullets",
        kicker="Урок 3",
        title="Вопрос задаётся до выхода, не на конусе",
        items=[
            "Есть ли лавинный комплект на группу.",
            "Умеют ли эти люди им пользоваться — не «когда-то слышали».",
            "Если ответа нет, это довод не выходить на лавинный рельеф.",
            "Этапы поиска называют на очном занятии. Этот слайд их не отрабатывает.",
        ],
        note="Если просят «хотя бы порядок действий», вернуть к границе модуля: имена этапов без приёма создают ложную готовность. Порядок отрабатывается с инструктором.",
    )
    d.add("section", number="УРОК 4", title="Холод, солнце, высота,\nпитание и вода", pages="Пособие, с. 122–139",
          note="Глава 7 — про влияние среды на человека. На слайдах нет диагнозов и назначений. Нормы и признаки сверять со страницами пособия.")
    d.add(
        "bullets",
        kicker="Урок 4",
        title="Среда снижает запас человека",
        items=[
            "Холод, солнце, вода, еда и высота меняют работоспособность и безопасность.",
            "Это данные для решения продолжать или разворачиваться, а не фон маршрута.",
            "Жалоба, плохой сон, сбой координации — информация, не помеха плану.",
            "Терпеть до вершины и скрывать состояние от группы — ошибочные действия.",
        ],
        note="Не переходить к схемам лечения. Граница первой помощи — отдельный модуль и очная практика.",
    )
    d.add(
        "split",
        kicker="Урок 4",
        title="Форма и акклиматизация — не одно и то же",
        left={"label": "ЧТО ЧАСТО ПУТАЮТ", "text": "Сильный и тренированный участник «перенесёт высоту». Уверенность подменяет план акклиматизации."},
        right={"label": "КАК В МОДУЛЕ", "text": "Личная реакция на высоту может отличаться от формы, от плана и от остальных. План из-за этого не отменяют."},
        note="Прямая формулировка вопроса 6.",
    )
    d.add(
        "task",
        kicker="Урок 4  ·  задание",
        title="Жалоба меняет план всей группы",
        beginner="Свяжите самочувствие участника с решением по маршруту: что пересматривается сразу, а что нужно дополнительно узнать.",
        instructor="Объясните группе, почему разворачивают и сильного участника. Назовите, какой признак делает продолжение неприемлемым.",
        note="Ответ модуля на ухудшение: оценить состояние, пересмотреть план, при необходимости запросить помощь. Не игнорировать и не скрывать.",
    )
    d.add(
        "section",
        number="КОНТРОЛЬ",
        title="Шесть вопросов\nпо ситуации",
        pages="Ответы записывают до ключа",
        note="Ситуация из итогового задания модуля. Не помогать выбором буквы.",
    )
    d.add(
        "bullets",
        kicker="Ситуация модуля",
        title="Знакомый склон на следующий день уже не знакомый",
        items=[
            "После снегопада и ветра группа планирует пересечь склон по известному маршруту.",
            "Один участник плохо спал и жалуется на самочувствие.",
            "Дальше шесть вопросов. Букву выбирают до обсуждения.",
            "Ключ на следующем блоке не листаем вместе с группой.",
        ],
        note="Если в бумажном комплекте у ситуации есть продолжение на странице 7, дочитать его вслух из PDF, не из слайда.",
    )
    questions = [
        (1, "Что определяет лавинный риск?",
         [("А", "Только крутизна и наличие снега"),
          ("Б", "Снег, погода, рельеф и воздействие на людей"),
          ("В", "Лавинный бюллетень по району, без оценки конкретного склона")]),
        (2, "Проход предыдущей группы по склону означает…",
         [("А", "Склон безопасен: по нему уже прошли"),
          ("Б", "Недостаточную информацию для гарантии"),
          ("В", "Маршрут можно не пересматривать")]),
        (3, "Как понимать «красные флаги»?",
         [("А", "Повод пересмотреть план"),
          ("Б", "Если флагов не видно, склон можно считать устойчивым"),
          ("В", "Повод ускориться, пока обстановка не изменилась")]),
        (4, "Может ли онлайн-тест подтвердить навык лавинного поиска?",
         [("А", "Да, если человек верно назвал этапы поиска"),
          ("Б", "Нет, нужна очная практическая проверка"),
          ("В", "Да, если тест модуля сдан выше порога")]),
        (5, "Как реагировать на ухудшение самочувствия на высоте?",
         [("А", "Игнорировать до вершины"),
          ("Б", "Оценить состояние и пересмотреть план, при необходимости запросить помощь"),
          ("В", "Скрыть состояние от группы")]),
        (6, "Почему план акклиматизации нельзя заменить уверенностью в форме?",
         [("А", "Личная реакция на высоту может отличаться"),
          ("Б", "Тренированность заменяет акклиматизацию"),
          ("В", "План нужен только тем, кто на высоте впервые")]),
    ]
    for num, prompt, options in questions:
        d.add("question", qn=num, prompt=prompt, options=options,
              note="Не комментировать до всех шести ответов. Верная буква есть только на слайде ключа.")
    d.add("key")
    d.add("result")
    d.build()
    print(OUT, "slides", len(d.slides))


if __name__ == "__main__":
    main()
